import openpyxl, sys, io
from collections import Counter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

file_path = r"C:\Users\said\Downloads\APPELS-DOFFRE-AVIS-DATTRIBUTION-2026-GAZON.xlsx"
wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)

print("ALL SHEETS IN WORKBOOK:")
print(wb.sheetnames)

for sheet_name in ["APPEL D'OFFRE 2026", "AVIS D'ATTRIBUTIONS 2026", "APPEL D'OFFRE 2024", "AVIS D'ATTRIBUTIONS 2024"]:
    if sheet_name not in wb.sheetnames:
        continue
    ws = wb[sheet_name]
    print(f"\n{'='*30} {sheet_name} {'='*30}")
    
    rows = list(ws.iter_rows(values_only=True))
    header_row = None
    header_idx = None
    data_rows = []
    
    for idx, r in enumerate(rows):
        # Header is usually row 2 or 3 where Col 1 is 'N°'
        if r and any(str(c).strip() in ['N°', 'N'] for c in r[:3] if c is not None):
            header_row = [str(c).strip() if c is not None else "" for c in r]
            header_idx = idx
            break
            
    if header_idx is not None:
        print(f"Header at row index {header_idx + 1}:")
        for i, col_name in enumerate(header_row):
            if col_name:
                print(f"  Col {i+1}: {col_name}")
                
        for idx in range(header_idx + 1, len(rows)):
            r = rows[idx]
            if r and r[0] is not None and str(r[0]).strip() not in ['', 'None']:
                data_rows.append(r)
                
        print(f"\nTotal populated data rows: {len(data_rows)}")
        
        # Analyze columns
        # In both sheets, check Col 3 and Col 5
        col3_vals = [str(r[2]).strip() for r in data_rows if len(r) > 2 and r[2] is not None]
        col5_vals = [str(r[4]).strip() for r in data_rows if len(r) > 4 and r[4] is not None]
        
        print(f"\nTop 20 most common values in Col 3 ({header_row[2]}):")
        for val, count in Counter(col3_vals).most_common(20):
            print(f"  ({count:3d}x) {val}")
            
        print(f"\nTop 25 most common values in Col 5 ({header_row[4]}):")
        for val, count in Counter(col5_vals).most_common(25):
            print(f"  ({count:3d}x) {val}")
            
        # Also check other relevant columns
        if len(header_row) > 7:
            col8_vals = [str(r[7]).strip() for r in data_rows if len(r) > 7 and r[7] is not None]
            print(f"\nTop 10 most common values in Col 8 ({header_row[7]}):")
            for val, count in Counter(col8_vals).most_common(10):
                print(f"  ({count:3d}x) {val}")
