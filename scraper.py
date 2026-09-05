"""
AlgerieMarches Scraper v3 — single login, single disconnect, keep session.

How it works:
  1. Try saved cookies first (skip login entirely if they work).
  2. If no cookies or expired: login once -> disconnect ONE competitor
     (promotes our session from temp to normal) -> scrape with that session.
  3. Save cookies to disk for next run.
  
  Never signout. Never drain the pool. Just disconnect one and use.
"""

import csv
import json
import logging
import os
import re
import sys
import time
import unicodedata
from copy import copy
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote

import gspread
import openpyxl
import requests
from dotenv import load_dotenv



# ── Paths ────────────────────────────────────────────────────────────────────
SCRIPT_DIR = Path(__file__).resolve().parent
STATE_DIR = SCRIPT_DIR / "state"
DATA_DIR = SCRIPT_DIR / "data"
COOKIE_FILE = STATE_DIR / "cookies.json"
LOG_FILE = SCRIPT_DIR / "scraper.log"

STATE_DIR.mkdir(exist_ok=True)
DATA_DIR.mkdir(exist_ok=True)

# ── Logging ──────────────────────────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    handlers=[
        logging.FileHandler(LOG_FILE, encoding="utf-8"),
        logging.StreamHandler(
            open(sys.stdout.fileno(), mode="w", encoding="utf-8", closefd=False)
        ),
    ],
)
log = logging.getLogger("am-scraper")

# ── Config ───────────────────────────────────────────────────────────────────
load_dotenv(SCRIPT_DIR / ".env")

EMAIL = os.getenv("AM_EMAIL", "direction@tapidor.com")
PASSWORD = os.getenv("AM_PASSWORD", "")

BASE_URL = "https://algeriemarches.com"
API_BASE = "https://api.algeriemarches.com/api"

CSRF_URL = f"{BASE_URL}/api/auth/csrf"
LOGIN_URL = f"{BASE_URL}/api/auth/callback/credentials"
SESSIONS_URL = f"{BASE_URL}/api/auth/active-sessions"
DISCONNECT_URL = f"{BASE_URL}/api/proxy/auth/disconnect-session"

# Scrape config
ANNONCE_TYPE = "avis-attribution"
PAGE_SIZE = 20
MAX_PAGES = 3
KEYWORDS_ENV = os.getenv("AM_KEYWORDS", "alimentation|fourniture")
KEYWORD_FILTER = re.compile(KEYWORDS_ENV, re.IGNORECASE) if KEYWORDS_ENV else None

UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36"
)

# ── Excel Helpers (Downloads master sync) ─────────────────────────────────────

def strip_accents(text: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn")

ACTIONS = [
    "REALISATION, REHABILITATION ET REVETEMENT",
    "REALISATION,REHABILITATION ET REVETEMENT",
    "REALISATION ET AMENAGEMENT",
    "AMENAGEMENT ET REVETEMENT",
    "REVETEMENT ET AMENAGEMENT",
    "FOURNITURE ET POSE",
    "FOURNITURE ET INSTALLATION",
    "ETUDE ET SUIVI",
    "FOURNITURE",
    "REALISATION",
    "AMENAGEMENT",
    "REVETEMENT",
    "REHABILITATION",
    "RENOVATION",
    "ACHEVEMENT",
    "ACQUISITION",
    "APPROVISIONNEMENT",
    "ALIMENTATION",
    "ENTRETIEN",
    "TRAVAUX",
]

def parse_titre_and_type(titre: str):
    t_clean = titre.strip()
    t_norm = strip_accents(t_clean).upper()
    
    action = "REALISATION"
    for a in ACTIONS:
        if a in t_norm:
            action = a
            break
            
    known_projects = [
        "TERRAINS SPORTIFS", "TERRAIN SPORTIF",
        "TERRAINS DE FOOTBALL", "TERRAIN DE FOOTBALL",
        "STADE DE PROXIMITE", "TERRAIN SPORTIF DE PROXIMITE",
        "STADE DE FOOTBALL", "STADE MUNICIPAL", "STADE MATICO", "STADE",
        "SALLES MULTI-SPORTS", "SALLE MULTI-SPORTS", "SALLE DE SPORT",
        "GAZON SYNTHETIQUE", "AIRES DE JEUX", "AIRE DE JEUX",
        "ECOLES PRIMAIRES", "ECOLE PRIMAIRE", "MATICO",
        "ALIMENTATION SCOLAIRE", "EQUIPEMENTS BUREAUTIQUE ET INFORMATIQUES",
        "PRODUITS RADIO PHARMACEUTIQUES"
    ]
    
    proj_type = None
    for kp in known_projects:
        if kp in t_norm:
            proj_type = kp
            break
            
    if not proj_type:
        rest = t_norm
        for a in ACTIONS:
            rest = re.sub(rf"^\b{re.escape(a)}\b", "", rest).strip()
        rest = re.sub(r"^(?:DES|DU|DE LA|DE L'|DE|D'|AU PROFIT DU|AU PROFIT DE LA|AU PROFIT DE|POUR)\s+", "", rest).strip()
        rest = re.split(r"\b(?:AU PROFIT|POUR LE COMPTE|DANS LA WILAYA|A LA WILAYA)\b", rest)[0].strip()
        proj_type = rest[:40].strip() if rest else "DIVERS"

    return action, proj_type

def parse_commune(titre: str, annonceur: str) -> str:
    combined = f"{annonceur} {titre}".upper()
    m = re.search(r"\b(?:COMMUNE D'|COMMUNE DE |APC D'|APC DE )([A-Z\s'\-]+?)(?:\bWILAYA|\bDAIRA|\bALIMENTATION|$)", combined)
    if m:
        c = m.group(1).strip()
        return c if len(c) < 30 else c[:30].strip()
    if "COMMUNE" in annonceur.upper():
        words = annonceur.upper().replace("COMMUNE", "").replace("D'", "").replace("DE", "").strip()
        if words:
            return words[:30].strip()
    return "/"

def parse_budget(montant_str: str):
    if not montant_str:
        return "/"
    clean = re.sub(r"[^\d.,]", "", str(montant_str)).replace(",", ".")
    try:
        val = float(clean)
        return int(val) if val.is_integer() else val
    except Exception:
        return "/"

def parse_delai(nbr_jours: int, description: str) -> str:
    if nbr_jours and int(nbr_jours) > 0:
        return f"{nbr_jours} JOURS"
    m = re.search(r"\b(\d+\s*(?:JOURS?|MOIS))\b", str(description).upper())
    if m:
        return m.group(1).strip()
    return "/"

def parse_date(date_val):
    if not date_val:
        return "/"
    if isinstance(date_val, datetime):
        return datetime(date_val.year, date_val.month, date_val.day)
    try:
        s = str(date_val)[:10]
        dt = datetime.strptime(s, "%Y-%m-%d")
        return datetime(dt.year, dt.month, dt.day)
    except Exception:
        return "/"

def get_latest_downloads_excel() -> Path | None:
    downloads = Path.home() / "Downloads"
    if not downloads.exists():
        return None
    files = [f for f in downloads.glob("*.xlsx") if not f.name.startswith("~$")]
    if not files:
        return None
    files.sort(key=lambda f: f.stat().st_mtime, reverse=True)
    return files[0]

def append_results_to_excel(results: list[dict], excel_path: Path | None = None) -> int:
    """
    Append new scraped results into the latest downloads .xlsx file matching the existing structure.
    Returns the number of new rows added.
    """
    if excel_path is None:
        excel_path = get_latest_downloads_excel()
        
    if not excel_path or not excel_path.exists():
        log.warning("No .xlsx file found in Downloads to update.")
        return 0

    log.info("Checking Downloads Excel file: %s", excel_path.name)
    try:
        wb = openpyxl.load_workbook(excel_path)
    except PermissionError:
        log.warning("File %s is currently open. Will save updated copy to Downloads.", excel_path.name)
        excel_path = excel_path.parent / f"{excel_path.stem}_updated.xlsx"
        if excel_path.exists():
            wb = openpyxl.load_workbook(excel_path)
        else:
            return 0
    except Exception as e:
        log.error("Failed to open workbook %s: %s", excel_path, e)
        return 0

    # Determine sheet
    sheet_name = None
    curr_year = datetime.now().year
    candidates = [
        f"AVIS D'ATTRIBUTIONS {curr_year}",
        f"AVIS D'ATTRIBUTION {curr_year}",
        "AVIS D'ATTRIBUTIONS 2026",
        "AVIS D'ATTRIBUTIONS",
        "AVIS D'ATTRIBUTION",
    ]
    for c in candidates:
        if c in wb.sheetnames:
            sheet_name = c
            break
            
    if not sheet_name:
        sheet_name = wb.sheetnames[0]
        log.info("Target sheet fallback to: %s", sheet_name)
    else:
        log.info("Target Excel sheet: %s", sheet_name)

    ws = wb[sheet_name]

    last_row_idx = 3
    last_num = 0
    existing_items = set()

    for r in range(4, ws.max_row + 1):
        val_num = ws.cell(row=r, column=1).value
        if val_num is not None:
            last_row_idx = r
            try:
                last_num = int(val_num)
            except Exception:
                pass
            attr_par = str(ws.cell(row=r, column=13).value or "").strip().upper()
            wilaya = str(ws.cell(row=r, column=9).value or "").strip().upper()
            t_proj = str(ws.cell(row=r, column=5).value or "").strip().upper()
            existing_items.add((attr_par, wilaya, t_proj))
        else:
            empty = True
            for check in range(r, min(r + 10, ws.max_row + 1)):
                if ws.cell(row=check, column=1).value is not None:
                    empty = False
                    break
            if empty:
                break

    log.info("Sheet currently has %d rows (last N°: %d)", last_num, last_num)
    ref_row_idx = last_row_idx if last_row_idx >= 4 else 4

    added_count = 0
    today_dt = datetime(datetime.now().year, datetime.now().month, datetime.now().day)

    for item in results:
        titre = item.get("titre", "")
        annonceur = item.get("annonceur", "")
        action, ptype = parse_titre_and_type(titre)
        attr_par = item.get("entreprise_concernee", "").strip()
        wilaya = (item.get("wilaya") or "").strip().upper()
        
        # Deduplication check
        dedup_key = (attr_par.upper(), wilaya, ptype.upper())
        if dedup_key in existing_items and attr_par != "":
            log.info("  Already exists in Excel: %s (%s)", attr_par[:30], wilaya)
            continue

        last_num += 1
        last_row_idx += 1
        added_count += 1
        existing_items.add(dedup_key)

        commune = parse_commune(titre, annonceur)
        budget = parse_budget(item.get("montant", ""))
        delai = parse_delai(item.get("nbr_jours", 0), item.get("description", ""))
        dt_parution = parse_date(item.get("date_parution", ""))
        dt_echeance = parse_date(item.get("date_echeance", ""))

        ann_upper = annonceur.upper()
        if "COMMUNE" in ann_upper:
            ann_val = "COMMUNE"
        elif "DJS" in ann_upper:
            ann_val = "DJS DE LA WILAYA"
        elif "DEP" in ann_upper:
            ann_val = "DEP DE LA WILAYA"
        else:
            ann_val = annonceur[:40] if annonceur else "/"

        row_values = [
            last_num,                                # Col 1: N°
            today_dt,                                # Col 2: DATE
            action,                                  # Col 3: TITRE D'AVIS D'ATTRIBUTION
            1,                                       # Col 4: Nombre de projet
            ptype,                                   # Col 5: TYPE DE PROJET
            dt_parution,                             # Col 6: DATE DE PARUTION
            dt_echeance,                             # Col 7: DATE D'ECHEANCE
            ann_val,                                 # Col 8: ANNONCEUR
            wilaya if wilaya else "/",               # Col 9: WILAYA
            commune,                                 # Col 10: COMMUNE
            dt_parution,                             # Col 11: DATE D'ATTRIBUTION
            budget,                                  # Col 12: BUDGET
            attr_par if attr_par else "/",           # Col 13: ATTRIBUTION PAR
            delai,                                   # Col 14: DELAI
            "/",                                     # Col 15: WILAYA 2
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=last_row_idx, column=col_idx, value=val)
            ref_cell = ws.cell(row=ref_row_idx, column=col_idx)
            if ref_cell.font:
                cell.font = copy(ref_cell.font)
            if ref_cell.border:
                cell.border = copy(ref_cell.border)
            if ref_cell.alignment:
                cell.alignment = copy(ref_cell.alignment)
            if ref_cell.number_format:
                cell.number_format = ref_cell.number_format

    if added_count > 0:
        ws.cell(row=1, column=4, value="=TODAY()")
        try:
            wb.save(excel_path)
            log.info("Successfully appended %d new rows to %s (sheet: %s)", added_count, excel_path.name, sheet_name)
        except PermissionError:
            alt_path = excel_path.parent / f"{excel_path.stem}_updated.xlsx"
            wb.save(alt_path)
            log.warning("Original file was locked. Saved %d rows to: %s", added_count, alt_path.name)
    else:
        log.info("No new rows needed for Excel (all already present)")

    return added_count


def append_results_to_gsheet(results: list[dict], sheet_id: str | None = None) -> int:
    """
    Append new scraped results into the Google Sheet.
    Works seamlessly locally (state/service_account.json) and in GitHub Actions (GCP_SERVICE_ACCOUNT_KEY).
    """
    sheet_id = sheet_id or os.getenv("GOOGLE_SHEET_ID", "1y5-OxNeL_z8hCUNNEVoh1nKZVsyBYrq5EJcgKBEt920")
    if not sheet_id:
        log.info("No GOOGLE_SHEET_ID configured -- skipping Google Sheet sync.")
        return 0

    sa_json_env = os.getenv("GCP_SERVICE_ACCOUNT_KEY")
    sa_file = SCRIPT_DIR / "state" / "service_account.json"

    gc = None
    if sa_json_env:
        try:
            creds_dict = json.loads(sa_json_env)
            gc = gspread.service_account_from_dict(creds_dict)
            log.info("Authenticated with Google via GCP_SERVICE_ACCOUNT_KEY env var")
        except Exception as e:
            log.error("Failed to authenticate with GCP_SERVICE_ACCOUNT_KEY env var: %s", e)
    elif sa_file.exists():
        try:
            gc = gspread.service_account(filename=str(sa_file))
            log.info("Authenticated with Google via state/service_account.json")
        except Exception as e:
            log.error("Failed to authenticate with service_account.json file: %s", e)

    if not gc:
        log.info("No Google service account credentials found -- skipping Google Sheet sync.")
        return 0

    try:
        sh = gc.open_by_key(sheet_id)
        log.info("Opened Google Sheet: %s", sh.title)
    except Exception as e:
        log.warning("Could not access Google Sheet '%s': %s", sheet_id, e)
        log.warning(">> Please make sure you clicked 'Share' on your Google Sheet and added:")
        log.warning(">> algeriemarches-bot@project-df86d806-9e9d-42be-a7d.iam.gserviceaccount.com as Editor!")
        return 0

    curr_year = datetime.now().year
    candidates = [
        f"AVIS D'ATTRIBUTIONS {curr_year}",
        f"AVIS D'ATTRIBUTION {curr_year}",
        "AVIS D'ATTRIBUTIONS 2026",
        "AVIS D'ATTRIBUTIONS",
        "AVIS D'ATTRIBUTION",
    ]
    target_ws = None
    for ws in sh.worksheets():
        if ws.title in candidates:
            target_ws = ws
            break

    if not target_ws:
        target_ws = sh.sheet1
        log.info("Google Sheet target worksheet fallback to: %s", target_ws.title)
    else:
        log.info("Target Google Sheet worksheet: %s", target_ws.title)

    try:
        all_values = target_ws.get_all_values()
    except Exception as e:
        log.error("Failed to read values from Google Sheet: %s", e)
        return 0

    last_num = 0
    existing_items = set()

    for idx, row in enumerate(all_values):
        if idx < 3:
            continue
        if row and len(row) > 0 and str(row[0]).strip():
            try:
                n = int(str(row[0]).strip())
                if n > last_num:
                    last_num = n
            except Exception:
                pass
            attr_par = row[12].strip().upper() if len(row) > 12 else ""
            wilaya = row[8].strip().upper() if len(row) > 8 else ""
            t_proj = row[4].strip().upper() if len(row) > 4 else ""
            existing_items.add((attr_par, wilaya, t_proj))

    log.info("Google Sheet currently has %d rows (last N°: %d)", len(all_values) - 3 if len(all_values) >= 3 else 0, last_num)

    rows_to_append = []
    today_str = datetime.now().strftime("%Y-%m-%d")

    for item in results:
        titre = item.get("titre", "")
        annonceur = item.get("annonceur", "")
        action, ptype = parse_titre_and_type(titre)
        attr_par = item.get("entreprise_concernee", "").strip()
        wilaya = (item.get("wilaya") or "").strip().upper()

        dedup_key = (attr_par.upper(), wilaya, ptype.upper())
        if dedup_key in existing_items and attr_par != "":
            log.info("  Already in Google Sheet: %s (%s)", attr_par[:30], wilaya)
            continue

        last_num += 1
        existing_items.add(dedup_key)

        commune = parse_commune(titre, annonceur)
        budget = parse_budget(item.get("montant", ""))
        delai = parse_delai(item.get("nbr_jours", 0), item.get("description", ""))
        dt_parution = str(item.get("date_parution", ""))[:10] or "/"
        dt_echeance = str(item.get("date_echeance", ""))[:10] or "/"

        ann_upper = annonceur.upper()
        if "COMMUNE" in ann_upper:
            ann_val = "COMMUNE"
        elif "DJS" in ann_upper:
            ann_val = "DJS DE LA WILAYA"
        elif "DEP" in ann_upper:
            ann_val = "DEP DE LA WILAYA"
        else:
            ann_val = annonceur[:40] if annonceur else "/"

        row_vals = [
            last_num,
            today_str,
            action,
            1,
            ptype,
            dt_parution,
            dt_echeance,
            ann_val,
            wilaya if wilaya else "/",
            commune,
            dt_parution,
            budget,
            attr_par if attr_par else "/",
            delai,
            "/",
        ]
        rows_to_append.append(row_vals)

    if rows_to_append:
        try:
            target_ws.append_rows(rows_to_append, value_input_option="USER_ENTERED")
            log.info("Successfully appended %d new rows to Google Sheet '%s'", len(rows_to_append), target_ws.title)
        except Exception as e:
            log.error("Failed to append rows to Google Sheet: %s", e)
            return 0
    else:
        log.info("No new rows needed for Google Sheet (all already present)")

    return len(rows_to_append)


class AlgerieMarchesScraper:

    def __init__(self):
        self.s = requests.Session()
        self.s.headers.update({
            "User-Agent": UA,
            "Accept": "application/json",
            "Origin": BASE_URL,
            "Referer": f"{BASE_URL}/login",
        })
        self._load_cookies()

    # ── JSON cookie persistence ──────────────────────────────────────────────

    def _save_cookies(self):
        cookies = []
        for c in self.s.cookies:
            cookies.append({
                "name": c.name, "value": c.value,
                "domain": c.domain, "path": c.path,
                "secure": c.secure, "expires": c.expires,
            })
        with open(COOKIE_FILE, "w", encoding="utf-8") as f:
            json.dump(cookies, f, indent=2)
        log.info("Saved %d cookies to disk", len(cookies))

    def _load_cookies(self):
        if not COOKIE_FILE.exists():
            log.info("No saved cookies")
            return
        try:
            with open(COOKIE_FILE, "r", encoding="utf-8") as f:
                cookies = json.load(f)
            for c in cookies:
                self.s.cookies.set(
                    c["name"], c["value"],
                    domain=c.get("domain", ""),
                    path=c.get("path", "/"),
                    secure=c.get("secure", True),
                )
            log.info("Loaded %d cookies from disk", len(cookies))
        except Exception as e:
            log.warning("Failed to load cookies: %s", e)

    def _has_auth_cookies(self) -> bool:
        names = {c.name for c in self.s.cookies}
        return "__Secure-authjs.session-token" in names and "connect.sid" in names

    # ── Auth ─────────────────────────────────────────────────────────────────

    def _get_csrf(self) -> str:
        r = self.s.get(CSRF_URL, timeout=30)
        r.raise_for_status()
        return r.json()["csrfToken"]

    def _login(self, csrf: str) -> bool:
        """Login with JSON body (matching n8n workflow). Don't follow redirects."""
        r = self.s.post(
            LOGIN_URL,
            json={
                "csrfToken": csrf,
                "identifier": EMAIL,
                "email": EMAIL,
                "password": PASSWORD,
                "callbackUrl": BASE_URL,
            },
            headers={"Content-Type": "application/json"},
            allow_redirects=False,
            timeout=40,
        )
        has_auth = self._has_auth_cookies()
        loc = r.headers.get("Location", "")
        log.info("Login: status=%d location=%s has_auth=%s", r.status_code, loc[:80], has_auth)
        return has_auth

    def _get_connect_sid_id(self) -> str:
        for c in self.s.cookies:
            if c.name == "connect.sid":
                raw = unquote(unquote(c.value))
                return raw.lstrip("s:").split(".")[0]
        return ""

    def _list_sessions(self) -> list[dict]:
        try:
            r = self.s.get(SESSIONS_URL, timeout=30)
            data = r.json()
            if isinstance(data, dict) and "sessions" in data:
                return data["sessions"]
            return data if isinstance(data, list) else []
        except Exception as e:
            log.warning("List sessions failed: %s", e)
            return []

    def _disconnect_one(self, session_id: str) -> dict:
        r = self.s.post(DISCONNECT_URL, json={"session_id": session_id}, timeout=30)
        try:
            return r.json()
        except Exception:
            return {"status": r.status_code, "text": r.text[:200]}

    def _promote_session(self) -> bool:
        """Promote NextAuth session token from temporary to normal."""
        try:
            csrf = self._get_csrf()
            r = self.s.post(
                f"{BASE_URL}/api/auth/session",
                json={"csrfToken": csrf, "data": {"sessionType": "normal"}},
                headers={"Content-Type": "application/json"},
                timeout=15,
            )
            data = r.json()
            is_normal = (
                data.get("sessionType") == "normal"
                or (data.get("user") or {}).get("sessionType") == "normal"
            )
            log.info("Session promotion: status=%d normal=%s", r.status_code, is_normal)
            return is_normal
        except Exception as e:
            log.warning("Session promotion failed: %s", e)
            return False

    # ── Session lifecycle ────────────────────────────────────────────────────

    def _test_detail_access(self) -> bool:
        """Test that our session can access a full detail page (not manage-sessions redirect)."""
        if not self._has_auth_cookies():
            return False
        try:
            r = self.s.get(f"{API_BASE}/ads/all",
                           params={"type": ANNONCE_TYPE, "avis": "true", "page": 1, "limit": 1},
                           timeout=20)
            items = r.json().get("data", []) if isinstance(r.json(), dict) else []
            if not items:
                return False
            slug = items[0].get("slug", "")
            if not slug:
                return False

            r = self.s.get(f"{BASE_URL}/annonces/{slug}",
                           headers={"Accept": "text/html"},
                           timeout=30, allow_redirects=True)
            html = r.text
            if "manage-sessions" in html[:5000]:
                log.info("Detail test: redirected to manage-sessions (session is temp)")
                return False
            if len(html) < 50000:
                log.info("Detail test: page too small (%d bytes)", len(html))
                return False
            log.info("Detail test: OK (%d bytes)", len(html))
            return True
        except Exception as e:
            log.warning("Detail test error: %s", e)
            return False

    def ensure_session(self):
        """
        Get a working session.
        1. Try saved cookies. If temporary, promote to normal.
        2. If expired/invalid: login once -> disconnect AT MOST ONE competitor -> promote to normal -> done.
        """
        if not PASSWORD:
            raise RuntimeError("No password. Edit .env and set AM_PASSWORD=...")

        # Try saved cookies
        if self._has_auth_cookies():
            log.info("Have saved auth cookies, testing...")
            try:
                r_sess = self.s.get(f"{BASE_URL}/api/auth/session", timeout=15)
                sess_data = r_sess.json()
                if sess_data and sess_data.get("user"):
                    curr_type = sess_data.get("sessionType") or sess_data.get("user", {}).get("sessionType")
                    if curr_type == "temporary":
                        log.info("Saved session is temporary, promoting to normal...")
                        self._promote_session()
            except Exception as e:
                log.warning("Session check error: %s", e)

            if self._test_detail_access():
                log.info("Saved session works!")
                self._save_cookies()
                return
            log.info("Saved cookies didn't work, need fresh login")

        # Fresh login
        self.s.cookies.clear()
        csrf = self._get_csrf()
        if not self._login(csrf):
            raise RuntimeError("Login failed -- check password in .env")

        # List sessions and inspect
        my_id = self._get_connect_sid_id()
        sessions = self._list_sessions()
        my_type = "?"
        competitors = []
        for s in sessions:
            sid = s.get("sessionId", "")
            if sid == my_id:
                my_type = s.get("type", "?")
            elif sid:
                competitors.append(sid)

        log.info("Sessions: %d total, my_id=%s type=%s, competitors=%d",
                 len(sessions), my_id[:10] if my_id else "?", my_type, len(competitors))

        # Disconnect AT MOST ONE competitor if we are in temporary state
        if competitors and my_type == "temp":
            target = competitors[0]
            result = self._disconnect_one(target)
            log.info("Disconnected %s -> %s", target[:10],
                     json.dumps(result, ensure_ascii=False)[:200])
        else:
            log.info("No competitor disconnect needed (my_type=%s, competitors=%d)", my_type, len(competitors))

        # Promote session from temporary to normal
        self._promote_session()

        # Save the promoted session
        self._save_cookies()

    # ── Scraping ─────────────────────────────────────────────────────────────

    def fetch_listings(self, page: int = 1) -> list[dict]:
        r = self.s.get(f"{API_BASE}/ads/all",
                       params={"type": ANNONCE_TYPE, "avis": "true", "page": page, "limit": PAGE_SIZE},
                       timeout=30)
        data = r.json()
        return data.get("data", []) if isinstance(data, dict) else []

    def fetch_detail(self, slug: str) -> dict:
        """Fetch detail page and extract ad data from RSC payload + rendered HTML."""
        url = f"{BASE_URL}/annonces/{slug}"
        r = self.s.get(url, headers={"Accept": "text/html"}, timeout=40, allow_redirects=True)
        html = r.text

        if "manage-sessions" in html[:5000]:
            return {"_error": "SESSION_BLOCKED", "_html_len": len(html)}

        result = {"_source": "html", "_html_len": len(html)}

        # 1. Extract ad JSON from RSC payload
        unescaped = html.replace('\\"', '"').replace('\\\\', '\\')
        idx = unescaped.find('"ad":{')
        if idx >= 0:
            start = idx + 5
            depth, end = 0, start
            for i, ch in enumerate(unescaped[start:min(len(unescaped), start + 40000)]):
                if ch == '{':
                    depth += 1
                elif ch == '}':
                    depth -= 1
                    if depth == 0:
                        end = start + i + 1
                        break
            try:
                ad = json.loads(unescaped[start:end])
                result["id_annonce"] = ad.get("id_annonce", "")
                result["titre_detail"] = ad.get("titre", "")
                result["date_parution"] = str(ad.get("date_parution", ""))[:10]
                result["meta_description"] = ad.get("meta_description", "")
                result["wilaya"] = (ad.get("wilaya") or {}).get("nom_fr", "")
                result["type_annonce"] = (ad.get("type_annonce") or {}).get("type_annonce_fr", "")
                result["sous_type"] = (ad.get("type_appel_offres") or {}).get("types_ao_fr", "")
                secteurs = ad.get("secteurs") or []
                result["secteur"] = ", ".join([s.get("secteur_fr", "") for s in secteurs if s.get("secteur_fr")])

                # Paywalled fields present in authenticated RSC payload:
                annonceurs = ad.get("annonceurs") or {}
                result["annonceur"] = annonceurs.get("annonceur") or ad.get("adresse_annonceur") or ""
                result["code_annonce"] = ad.get("code_annonce") or ""
                result["date_echeance"] = str(ad.get("date_echeance", ""))[:10] if ad.get("date_echeance") else ""
                result["entreprise_concernee"] = ad.get("entreprise_concerne") or ""
                result["anep"] = ad.get("anep") or ""
                result["parution_media"] = (ad.get("parution_media") or {}).get("media", "")
                result["description"] = ad.get("description") or ""

                if ad.get("image_principale"):
                    result["image_principale"] = f"{API_BASE}/images/{ad['image_principale']}"
                if ad.get("image_secondaire"):
                    result["image_secondaire"] = f"{API_BASE}/images/{ad['image_secondaire']}"

            except Exception as e:
                log.warning("Failed to parse ad JSON for %s: %s", slug[:30], e)

        # 2. Extract rendered HTML fallback / additional fields
        m = re.search(r'Annonceur</span>.*?capitalize[^>]*>([^<]+)<', unescaped, re.DOTALL)
        if m and m.group(1).strip() and not result.get("annonceur"):
            result["annonceur"] = m.group(1).strip()

        m = re.search(r'Code\s*N[^:]*:\s*(\S+)', unescaped)
        if m and not result.get("code_annonce"):
            result["code_annonce"] = m.group(1).strip()

        m = re.search(r"Date d.{1,5}ch[eé]ance.*?(\d{4}-\d{2}-\d{2})", unescaped, re.DOTALL)
        if m and not result.get("date_echeance"):
            result["date_echeance"] = m.group(1)

        m = re.search(r'Entreprise concern[eé]e.*?</span>\s*<[^>]*>([^<]+)<', unescaped, re.DOTALL | re.IGNORECASE)
        if m and m.group(1).strip() and not result.get("entreprise_concernee"):
            result["entreprise_concernee"] = m.group(1).strip()

        m = re.search(r'Objet du projet\s*:?\s*</[^>]+>\s*([^<]+)', unescaped, re.DOTALL)
        if m and m.group(1).strip():
            result["objet"] = m.group(1).strip()

        # Montant extraction (if present in text/description)
        combined_text = f"{result.get('description', '')} {result.get('entreprise_concernee', '')}"
        m_montant = re.search(r'(\d[\d\s,.]*\s*(?:DA|DZD|dinars?))\b', combined_text, re.IGNORECASE)
        if m_montant:
            result["montant"] = m_montant.group(1).strip()
        else:
            result["montant"] = result.get("montant", "")

        # Subscription status
        if '"code":"RESTRICTED"' in html or '\\"code\\":\\"RESTRICTED\\"' in html:
            result["_subscription"] = "RESTRICTED"
        elif '"code":"ACTIVE"' in html or '\\"code\\":\\"ACTIVE\\"' in html or result.get("code_annonce"):
            result["_subscription"] = "ACTIVE"
        else:
            result["_subscription"] = "unknown"

        return result

    # ── Main ─────────────────────────────────────────────────────────────────

    def run(self):
        run_id = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        log.info("=" * 60)
        log.info("Run %s started", run_id)

        self.ensure_session()

        # Fetch listings
        all_annonces = []
        for page in range(1, MAX_PAGES + 1):
            listings = self.fetch_listings(page)
            if not listings:
                log.info("Page %d: 0 results -- stopping", page)
                break
            log.info("Page %d: %d annonces", page, len(listings))
            all_annonces.extend(listings)
            time.sleep(0.5)

        log.info("Total: %d annonces", len(all_annonces))

        filtered = [a for a in all_annonces if KEYWORD_FILTER.search(a.get("titre", ""))]
        log.info("After filter: %d annonces", len(filtered))

        if not filtered:
            log.info("No matches -- done")
            return

        # Fetch details
        results = []
        errors = []
        for a in filtered:
            ann_id = str(a.get("id_annonce", ""))
            slug = a.get("slug", "")
            titre = a.get("titre", "")
            raw_date = str(a.get("date_parution", ""))[:10]
            parts = raw_date.split("-")
            date_fr = f"{parts[2]}-{parts[1]}-{parts[0]}" if len(parts) == 3 else ""

            detail = None
            for attempt in range(3):
                try:
                    detail = self.fetch_detail(slug) if slug else {}
                    if detail and not detail.get("_error"):
                        break
                except requests.exceptions.RequestException as e:
                    wait = 5 * (attempt + 1)
                    log.warning("  Attempt %d/3 for %s: %s -- retry in %ds",
                                attempt + 1, ann_id, type(e).__name__, wait)
                    time.sleep(wait)

            if not detail or detail.get("_error"):
                detail = detail or {}
                errors.append(ann_id)
                log.warning("  SKIP %s", ann_id)

            row = {
                "run_id": run_id,
                "id": ann_id,
                "slug": slug,
                "url": f"{BASE_URL}/annonces/{slug}",
                "titre": titre,
                "wilaya": detail.get("wilaya") or (a.get("wilaya") or {}).get("nom_fr", ""),
                "date_parution_fr": date_fr,
                "date_echeance": detail.get("date_echeance", ""),
                "type_annonce": detail.get("type_annonce") or (a.get("type_annonce") or {}).get("type_annonce_fr", ""),
                "sous_type": detail.get("sous_type") or (a.get("type_appel_offres") or {}).get("types_ao_fr", ""),
                "secteur": detail.get("secteur") or ((a.get("secteurs") or [{}])[0].get("secteur_fr", "") if a.get("secteurs") else ""),
                "annonceur": detail.get("annonceur", ""),
                "code_annonce": detail.get("code_annonce", ""),
                "entreprise_concernee": detail.get("entreprise_concernee", ""),
                "montant": detail.get("montant", ""),
                "anep": detail.get("anep", ""),
                "parution_media": detail.get("parution_media", ""),
                "image_principale": detail.get("image_principale", ""),
                "image_secondaire": detail.get("image_secondaire", ""),
                "objet": detail.get("objet", ""),
                "description": detail.get("description", ""),
                "nbr_jours": detail.get("nbr_jours", 0),
                "date_parution": detail.get("date_parution") or str(a.get("date_parution", ""))[:10],
                "subscription": detail.get("_subscription", ""),
                "scraped_at": datetime.now().isoformat(),
            }
            results.append(row)
            log.info("  [OK] %s -- %s | Annonceur: %s | Code: %s (sub=%s)",
                     ann_id, titre[:40], detail.get("annonceur", "N/A")[:30],
                     detail.get("code_annonce", "N/A"), detail.get("_subscription", "?"))
            time.sleep(0.8)

        if errors:
            log.warning("Failed: %s", errors)

        # Write CSV
        csv_path = DATA_DIR / f"annonces_{run_id}.csv"
        if results:
            with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.DictWriter(f, fieldnames=results[0].keys(), delimiter=";")
                w.writeheader()
                w.writerows(results)
            log.info("Wrote %d rows -> %s", len(results), csv_path)

        master_csv = DATA_DIR / "annonces_master.csv"
        exists = master_csv.exists()
        if results:
            with open(master_csv, "a", newline="", encoding="utf-8-sig") as f:
                w = csv.DictWriter(f, fieldnames=results[0].keys(), delimiter=";")
                if not exists:
                    w.writeheader()
                w.writerows(results)
            log.info("Appended to master -> %s", master_csv)

        # Append directly to the latest .xlsx in Downloads
        if results:
            try:
                added_excel = append_results_to_excel(results)
                log.info("Excel sync complete: %d rows added", added_excel)
            except Exception as e:
                log.warning("Failed to sync to Excel: %s", e)

            # Append directly to Google Sheet
            try:
                added_gsheet = append_results_to_gsheet(results)
                log.info("Google Sheet sync complete: %d rows added", added_gsheet)
            except Exception as e:
                log.warning("Failed to sync to Google Sheet: %s", e)

        self._save_cookies()
        log.info("Run %s done -- %d results", run_id, len(results))


if __name__ == "__main__":
    if not PASSWORD:
        print("No password. Edit .env: AM_PASSWORD=your_password")
        sys.exit(1)
    scraper = AlgerieMarchesScraper()
    try:
        scraper.run()
    except Exception:
        log.exception("Fatal error")
        sys.exit(1)
